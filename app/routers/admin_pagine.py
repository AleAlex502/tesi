import os
import io
import uuid
from typing import Optional, List

from fastapi import APIRouter, Request, Form, HTTPException, UploadFile, File
from fastapi.responses import RedirectResponse
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
import mysql.connector

from app.database import get_connection
from app.auth import verifica_password, crea_access_token, get_admin_by_username, admin_da_token, hash_password
from app import crud

router = APIRouter(prefix="/admin", tags=["Pannello amministrazione"], include_in_schema=False)
templates = Jinja2Templates(directory="app/templates")

NOME_COOKIE = "access_token"
CARTELLA_ALLEGATI = "app/static/allegati"
DIMENSIONE_MASSIMA_PDF = 10 * 1024 * 1024  # 10 MB


def _salva_allegato_pdf(file: UploadFile) -> Optional[str]:
    """
    Salva il PDF caricato su disco con un nome univoco e restituisce il nome
    del file salvato (da mettere nella colonna allegato_pdf), o None se
    non è stato caricato nessun file valido.
    """
    if not file or not file.filename:
        return None

    if file.content_type != "application/pdf":
        raise HTTPException(status_code=400, detail="L'allegato deve essere un file PDF")

    contenuto = file.file.read()
    if len(contenuto) > DIMENSIONE_MASSIMA_PDF:
        raise HTTPException(status_code=400, detail="Il PDF supera la dimensione massima di 10 MB")

    os.makedirs(CARTELLA_ALLEGATI, exist_ok=True)
    nome_file = f"{uuid.uuid4().hex}.pdf"
    with open(os.path.join(CARTELLA_ALLEGATI, nome_file), "wb") as f:
        f.write(contenuto)

    return nome_file


def _elimina_allegato_pdf(nome_file: Optional[str]):
    if not nome_file:
        return
    percorso = os.path.join(CARTELLA_ALLEGATI, nome_file)
    if os.path.exists(percorso):
        os.remove(percorso)


def _admin_autenticato(request: Request):
    """Restituisce l'admin loggato leggendo il cookie, o None se non è loggato."""
    token = request.cookies.get(NOME_COOKIE)
    if not token:
        return None
    return admin_da_token(token)


def _int_o_none(valore: Optional[str]):
    if valore and valore.strip().lstrip("-").isdigit():
        return int(valore.strip())
    return None


def _testo_o_none(valore: Optional[str]):
    valore = (valore or "").strip()
    return valore or None


# ---------- Login / logout ----------

@router.get("/login")
def form_login(request: Request, errore: Optional[str] = None):
    if _admin_autenticato(request):
        return RedirectResponse("/admin", status_code=303)
    return templates.TemplateResponse("admin_login.html", {"request": request, "errore": errore})


@router.post("/login")
def esegui_login(username: str = Form(...), password: str = Form(...)):
    admin = get_admin_by_username(username)

    if not admin or not verifica_password(password, admin["password_hash"]):
        return RedirectResponse("/admin/login?errore=1", status_code=303)

    token = crea_access_token(admin["username"])
    risposta = RedirectResponse("/admin", status_code=303)
    risposta.set_cookie(
        key=NOME_COOKIE, value=token,
        httponly=True, samesite="lax", max_age=60 * 60 * 2  # 2 ore, come il token
    )
    return risposta


@router.get("/logout")
def logout():
    risposta = RedirectResponse("/admin/login", status_code=303)
    risposta.delete_cookie(NOME_COOKIE)
    return risposta


# ---------- Dashboard ----------

@router.get("")
def dashboard(request: Request, q: Optional[str] = None, esito: Optional[str] = None):
    admin = _admin_autenticato(request)
    if not admin:
        return RedirectResponse("/admin/login", status_code=303)

    connessione = get_connection()
    try:
        cursor = connessione.cursor(dictionary=True)
        if q:
            cursor.execute(
                "SELECT id_materiale, titolo, anno, tipologia FROM Materiale WHERE titolo LIKE %s ORDER BY titolo",
                (f"%{q}%",)
            )
        else:
            cursor.execute("SELECT id_materiale, titolo, anno, tipologia FROM Materiale ORDER BY titolo")
        materiali = cursor.fetchall()
        cursor.close()
    finally:
        connessione.close()

    return templates.TemplateResponse(
        "admin_dashboard.html",
        {"request": request, "admin": admin, "materiali": materiali, "q": q or "", "esito": esito}
    )


# ---------- Campi comuni al form di creazione e modifica ----------

def _leggi_campi_form(
    titolo, anno, num_pagine, isbn, tipologia,
    editore_nome, editore_citta, lingua_testo, lingua_originale,
    ebp, livello, link_catalogo, abstract, indice,
):
    return {
        "titolo": titolo.strip(),
        "anno": _int_o_none(anno),
        "num_pagine": _int_o_none(num_pagine),
        "isbn": _testo_o_none(isbn),
        "abstract": _testo_o_none(abstract),
        "indice": _testo_o_none(indice),
        "tipologia": _testo_o_none(tipologia),
        "ebp": _testo_o_none(ebp),
        "livello": _testo_o_none(livello),
        "link_catalogo": _testo_o_none(link_catalogo),
        "editore_nome": _testo_o_none(editore_nome),
        "editore_citta": _testo_o_none(editore_citta),
        "lingua_testo": _testo_o_none(lingua_testo),
        "lingua_originale": _testo_o_none(lingua_originale),
    }


# ---------- Creazione ----------

@router.get("/nuovo")
def form_nuovo(request: Request):
    admin = _admin_autenticato(request)
    if not admin:
        return RedirectResponse("/admin/login", status_code=303)

    return templates.TemplateResponse(
        "admin_form.html",
        {
            "request": request, "admin": admin,
            "titolo_pagina": "Nuovo materiale", "azione": "/admin/nuovo",
            "valori": {}, "persone_testo": "", "copie_testo": "", "tag_testo": "",
            "sh_testo": "", "categoria_testo": "",
        }
    )


@router.post("/nuovo")
def crea_nuovo(
    request: Request,
    titolo: str = Form(...),
    anno: Optional[str] = Form(None),
    num_pagine: Optional[str] = Form(None),
    isbn: Optional[str] = Form(None),
    tipologia: Optional[str] = Form(None),
    editore_nome: Optional[str] = Form(None),
    editore_citta: Optional[str] = Form(None),
    lingua_testo: Optional[str] = Form(None),
    lingua_originale: Optional[str] = Form(None),
    ebp: Optional[str] = Form(None),
    livello: Optional[str] = Form(None),
    link_catalogo: Optional[str] = Form(None),
    abstract: Optional[str] = Form(None),
    indice: Optional[str] = Form(None),
    categoria: Optional[str] = Form(None),
    tag: Optional[str] = Form(None),
    subject_headings: Optional[str] = Form(None),
    persone: Optional[str] = Form(None),
    copie: Optional[str] = Form(None),
    allegato_pdf: Optional[UploadFile] = File(None),
):
    admin = _admin_autenticato(request)
    if not admin:
        return RedirectResponse("/admin/login", status_code=303)

    campi = _leggi_campi_form(
        titolo, anno, num_pagine, isbn, tipologia,
        editore_nome, editore_citta, lingua_testo, lingua_originale,
        ebp, livello, link_catalogo, abstract, indice,
    )
    campi["allegato_pdf"] = _salva_allegato_pdf(allegato_pdf)

    connessione = get_connection()
    try:
        id_materiale = crud.crea_materiale_completo(
            connessione, campi,
            crud.parse_persone_testo(persone),
            _testo_o_none(categoria),
            crud.parse_lista_semplice(tag),
            crud.parse_lista_semplice(subject_headings),
            crud.parse_copie_testo(copie),
        )
    except mysql.connector.IntegrityError as errore:
        if "isbn" in str(errore).lower():
            messaggio_errore = f"Esiste già un materiale con ISBN \"{isbn}\". Modifica l'ISBN o cerca il materiale già presente nel catalogo."
        else:
            messaggio_errore = f"Impossibile salvare: {errore}"

        return templates.TemplateResponse(
            "admin_form.html",
            {
                "request": request, "admin": admin,
                "titolo_pagina": "Nuovo materiale", "azione": "/admin/nuovo",
                "valori": campi,
                "persone_testo": persone or "", "copie_testo": copie or "",
                "tag_testo": tag or "", "sh_testo": subject_headings or "",
                "categoria_testo": categoria or "",
                "errore": messaggio_errore,
            }
        )
    finally:
        connessione.close()

    return RedirectResponse(f"/admin?esito=creato-{id_materiale}", status_code=303)


# ---------- Modifica ----------

def _carica_materiale_completo(connessione, id_materiale):
    cursor = connessione.cursor(dictionary=True)
    cursor.execute("SELECT * FROM Materiale WHERE id_materiale = %s", (id_materiale,))
    materiale = cursor.fetchone()

    if not materiale:
        cursor.close()
        return None

    cursor.execute(
        """
        SELECT p.nome, p.cognome, mp.ruolo
        FROM Materiale_Persona mp JOIN Persona p ON mp.id_persona = p.id_persona
        WHERE mp.id_materiale = %s
        """,
        (id_materiale,)
    )
    materiale["persone"] = cursor.fetchall()

    cursor.execute(
        "SELECT c.nome FROM Materiale_Categoria mc JOIN Categoria c ON mc.id_categoria = c.id_categoria WHERE mc.id_materiale = %s",
        (id_materiale,)
    )
    materiale["categorie"] = [r["nome"] for r in cursor.fetchall()]

    cursor.execute(
        "SELECT t.nome FROM Materiale_Tag mt JOIN Tag t ON mt.id_tag = t.id_tag WHERE mt.id_materiale = %s",
        (id_materiale,)
    )
    materiale["tag"] = [r["nome"] for r in cursor.fetchall()]

    cursor.execute(
        "SELECT sh.termine FROM Materiale_SH msh JOIN SH sh ON msh.id_sh = sh.id_sh WHERE msh.id_materiale = %s",
        (id_materiale,)
    )
    materiale["subject_headings"] = [r["termine"] for r in cursor.fetchall()]

    cursor.execute(
        "SELECT numero, collocazione, stato, note, nota_critica FROM Copia WHERE id_materiale = %s",
        (id_materiale,)
    )
    materiale["copie"] = cursor.fetchall()

    cursor.execute(
        "SELECT nome AS editore_nome, citta AS editore_citta FROM Editore WHERE id_editore = %s",
        (materiale.get("id_editore"),)
    )
    editore = cursor.fetchone() or {}

    cursor.execute("SELECT nome FROM Lingua WHERE id_lingua = %s", (materiale.get("id_lingua_testo"),))
    r = cursor.fetchone()
    lingua_testo = r["nome"] if r else None

    cursor.execute("SELECT nome FROM Lingua WHERE id_lingua = %s", (materiale.get("id_lingua_originale"),))
    r = cursor.fetchone()
    lingua_originale = r["nome"] if r else None

    cursor.close()

    materiale["editore_nome"] = editore.get("editore_nome")
    materiale["editore_citta"] = editore.get("editore_citta")
    materiale["lingua_testo"] = lingua_testo
    materiale["lingua_originale"] = lingua_originale

    return materiale


@router.get("/modifica/{id_materiale}")
def form_modifica(request: Request, id_materiale: int):
    admin = _admin_autenticato(request)
    if not admin:
        return RedirectResponse("/admin/login", status_code=303)

    connessione = get_connection()
    try:
        materiale = _carica_materiale_completo(connessione, id_materiale)
    finally:
        connessione.close()

    if not materiale:
        raise HTTPException(status_code=404, detail="Materiale non trovato")

    persone_testo = "\n".join(
        f"{p['nome']} {p['cognome']} ({p['ruolo']})".strip() for p in materiale["persone"]
    )
    copie_testo = "\n".join(
        f"{c['numero']};{c['collocazione'] or ''};{c['stato'] or ''};{c['nota_critica'] or c['note'] or ''}"
        for c in materiale["copie"]
    )

    return templates.TemplateResponse(
        "admin_form.html",
        {
            "request": request, "admin": admin,
            "titolo_pagina": f"Modifica: {materiale['titolo']}",
            "azione": f"/admin/modifica/{id_materiale}",
            "valori": materiale,
            "persone_testo": persone_testo,
            "copie_testo": copie_testo,
            "tag_testo": ", ".join(materiale["tag"]),
            "sh_testo": ", ".join(materiale["subject_headings"]),
            "categoria_testo": materiale["categorie"][0] if materiale["categorie"] else "",
        }
    )


@router.post("/modifica/{id_materiale}")
def salva_modifica(
    request: Request,
    id_materiale: int,
    titolo: str = Form(...),
    anno: Optional[str] = Form(None),
    num_pagine: Optional[str] = Form(None),
    isbn: Optional[str] = Form(None),
    tipologia: Optional[str] = Form(None),
    editore_nome: Optional[str] = Form(None),
    editore_citta: Optional[str] = Form(None),
    lingua_testo: Optional[str] = Form(None),
    lingua_originale: Optional[str] = Form(None),
    ebp: Optional[str] = Form(None),
    livello: Optional[str] = Form(None),
    link_catalogo: Optional[str] = Form(None),
    abstract: Optional[str] = Form(None),
    indice: Optional[str] = Form(None),
    categoria: Optional[str] = Form(None),
    tag: Optional[str] = Form(None),
    subject_headings: Optional[str] = Form(None),
    persone: Optional[str] = Form(None),
    copie: Optional[str] = Form(None),
    allegato_pdf: Optional[UploadFile] = File(None),
    rimuovi_allegato: Optional[str] = Form(None),
):
    admin = _admin_autenticato(request)
    if not admin:
        return RedirectResponse("/admin/login", status_code=303)

    campi = _leggi_campi_form(
        titolo, anno, num_pagine, isbn, tipologia,
        editore_nome, editore_citta, lingua_testo, lingua_originale,
        ebp, livello, link_catalogo, abstract, indice,
    )

    connessione = get_connection()
    try:
        cursor = connessione.cursor()
        cursor.execute("SELECT allegato_pdf FROM Materiale WHERE id_materiale = %s", (id_materiale,))
        riga = cursor.fetchone()
        cursor.close()

        if riga is None:
            raise HTTPException(status_code=404, detail="Materiale non trovato")

        allegato_attuale = riga[0]
        nuovo_file = _salva_allegato_pdf(allegato_pdf)

        if nuovo_file:
            # è stato caricato un nuovo PDF: sostituisce quello vecchio
            _elimina_allegato_pdf(allegato_attuale)
            campi["allegato_pdf"] = nuovo_file
        elif rimuovi_allegato:
            # l'admin ha spuntato "rimuovi allegato" senza caricarne uno nuovo
            _elimina_allegato_pdf(allegato_attuale)
            campi["allegato_pdf"] = None
        else:
            # nessuna modifica all'allegato: resta quello che c'era
            campi["allegato_pdf"] = allegato_attuale

        crud.aggiorna_materiale_completo(
            connessione, id_materiale, campi,
            crud.parse_persone_testo(persone),
            _testo_o_none(categoria),
            crud.parse_lista_semplice(tag),
            crud.parse_lista_semplice(subject_headings),
            crud.parse_copie_testo(copie),
        )
    except mysql.connector.IntegrityError as errore:
        if "isbn" in str(errore).lower():
            messaggio_errore = f"Esiste già un altro materiale con ISBN \"{isbn}\". Modifica l'ISBN per continuare."
        else:
            messaggio_errore = f"Impossibile salvare: {errore}"

        return templates.TemplateResponse(
            "admin_form.html",
            {
                "request": request, "admin": admin,
                "titolo_pagina": f"Modifica: {titolo}",
                "azione": f"/admin/modifica/{id_materiale}",
                "valori": campi,
                "persone_testo": persone or "", "copie_testo": copie or "",
                "tag_testo": tag or "", "sh_testo": subject_headings or "",
                "categoria_testo": categoria or "",
                "errore": messaggio_errore,
            }
        )
    finally:
        connessione.close()

    return RedirectResponse(f"/admin?esito=modificato-{id_materiale}", status_code=303)


# ---------- Eliminazione ----------

@router.post("/elimina/{id_materiale}")
def elimina(request: Request, id_materiale: int):
    admin = _admin_autenticato(request)
    if not admin:
        return RedirectResponse("/admin/login", status_code=303)
    if admin.get("ruolo") != "Amministratore":
        raise HTTPException(status_code=403, detail="Operazione riservata agli amministratori")

    connessione = get_connection()
    try:
        allegato = crud.elimina_materiale(connessione, id_materiale)
    finally:
        connessione.close()
    _elimina_allegato_pdf(allegato)

    return RedirectResponse("/admin?esito=eliminato", status_code=303)


@router.post("/elimina-multipli")
def elimina_multipli(request: Request, id_selezionati: List[int] = Form(...)):
    admin = _admin_autenticato(request)
    if not admin:
        return RedirectResponse("/admin/login", status_code=303)
    if admin.get("ruolo") != "Amministratore":
        raise HTTPException(status_code=403, detail="Operazione riservata agli amministratori")

    connessione = get_connection()
    try:
        for id_materiale in id_selezionati:
            allegato = crud.elimina_materiale(connessione, id_materiale)
            _elimina_allegato_pdf(allegato)
    finally:
        connessione.close()

    return RedirectResponse(f"/admin?esito=eliminati-{len(id_selezionati)}", status_code=303)


# ---------- Gestione utenti (solo Amministratore) ----------

def _richiedi_amministratore_html(request: Request):
    """
    Indica se l'utente è autenticato come amministratore.
    """
    admin = _admin_autenticato(request)
    if not admin:
        return None, RedirectResponse("/admin/login", status_code=303)
    if admin.get("ruolo") != "Amministratore":
        return None, RedirectResponse("/admin?esito=permessi-insufficienti", status_code=303)
    return admin, None

def _conta_amministratori(connessione) -> int:
    cursor = connessione.cursor()
    cursor.execute("SELECT COUNT(*) FROM Amministratore WHERE ruolo = 'Amministratore'")
    n = cursor.fetchone()[0]
    cursor.close()
    return n
 

@router.get("/utenti")
def elenco_utenti(request: Request):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect

    connessione = get_connection()
    try:
        cursor = connessione.cursor(dictionary=True)
        cursor.execute("SELECT id_admin, username, ruolo FROM Amministratore ORDER BY username")
        utenti = cursor.fetchall()
        cursor.close()
    finally:
        connessione.close()

    return templates.TemplateResponse(
        "admin_utenti.html", {"request": request, "admin": admin, "utenti": utenti}
    )


@router.get("/utenti/nuovo")
def form_nuovo_utente(request: Request):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect

    return templates.TemplateResponse(
        "admin_utente_form.html", {"request": request, "admin": admin, "errore": None}
    )


@router.post("/utenti/nuovo")
def crea_nuovo_utente(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    conferma_password: str = Form(...),
    ruolo: str = Form(...),
):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect

    username = username.strip()
    ruolo = ruolo if ruolo in ("Bibliotecario", "Amministratore") else "Bibliotecario"

    errore = None
    if password != conferma_password:
        errore = "Le due password non coincidono."
    elif len(password) < 8:
        errore = "La password deve avere almeno 8 caratteri."

    if not errore:
        connessione = get_connection()
        try:
            cursor = connessione.cursor()
            cursor.execute("SELECT id_admin FROM Amministratore WHERE username = %s", (username,))
            if cursor.fetchone():
                errore = f"Esiste già un utente con username \"{username}\"."
            cursor.close()

            if not errore:
                cursor = connessione.cursor()
                cursor.execute(
                    "INSERT INTO Amministratore (username, password_hash, ruolo) VALUES (%s, %s, %s)",
                    (username, hash_password(password), ruolo)
                )
                connessione.commit()
                cursor.close()
        finally:
            connessione.close()

    if errore:
        return templates.TemplateResponse(
            "admin_utente_form.html",
            {"request": request, "admin": admin, "errore": errore}
        )

    return RedirectResponse("/admin/utenti?esito=creato", status_code=303)

# ---------- Reimposta password di un utente (solo Amministratore) ----------
 
@router.get("/utenti/{id_admin}/reimposta-password")
def form_reimposta_password(request: Request, id_admin: int):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect
 
    connessione = get_connection()
    try:
        cursor = connessione.cursor(dictionary=True)
        cursor.execute("SELECT id_admin, username FROM Amministratore WHERE id_admin = %s", (id_admin,))
        utente = cursor.fetchone()
        cursor.close()
    finally:
        connessione.close()
 
    if not utente:
        raise HTTPException(status_code=404, detail="Utente non trovato")
 
    return templates.TemplateResponse(
        "admin_reimposta_password.html",
        {"request": request, "admin": admin, "utente": utente, "errore": None}
    )
 

@router.get("/utenti/{id_admin}/modifica")
def form_modifica_utente(request: Request, id_admin: int):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect
 
    connessione = get_connection()
    try:
        cursor = connessione.cursor(dictionary=True)
        cursor.execute("SELECT id_admin, username, ruolo FROM Amministratore WHERE id_admin = %s", (id_admin,))
        utente = cursor.fetchone()
        cursor.close()
    finally:
        connessione.close()
 
    if not utente:
        raise HTTPException(status_code=404, detail="Utente non trovato")
 
    return templates.TemplateResponse(
        "admin_utente_modifica.html",
        {"request": request, "admin": admin, "utente": utente, "errore": None}
    )
 
 
@router.post("/utenti/{id_admin}/modifica")
def salva_modifica_utente(
    request: Request,
    id_admin: int,
    username: str = Form(...),
    ruolo: str = Form(...),
):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect
 
    username = username.strip()
    ruolo = ruolo if ruolo in ("Bibliotecario", "Amministratore") else "Bibliotecario"
 
    connessione = get_connection()
    try:
        cursor = connessione.cursor(dictionary=True)
        cursor.execute("SELECT id_admin, username, ruolo FROM Amministratore WHERE id_admin = %s", (id_admin,))
        utente = cursor.fetchone()
        cursor.close()
 
        if not utente:
            raise HTTPException(status_code=404, detail="Utente non trovato")
 
        errore = None
        if not username:
            errore = "Lo username non può essere vuoto."
        elif (
            ruolo != "Amministratore"
            and utente["ruolo"] == "Amministratore"
            and _conta_amministratori(connessione) <= 1
        ):
            errore = "Non puoi togliere il ruolo di Amministratore all'unico account amministratore rimasto."
 
        if not errore:
            cursor = connessione.cursor()
            cursor.execute(
                "SELECT id_admin FROM Amministratore WHERE username = %s AND id_admin != %s",
                (username, id_admin)
            )
            if cursor.fetchone():
                errore = f"Esiste già un altro utente con username \"{username}\"."
            cursor.close()
 
        if errore:
            return templates.TemplateResponse(
                "admin_utente_modifica.html",
                {
                    "request": request, "admin": admin,
                    "utente": {"id_admin": id_admin, "username": username, "ruolo": ruolo},
                    "errore": errore,
                }
            )
 
        cursor = connessione.cursor()
        cursor.execute(
            "UPDATE Amministratore SET username = %s, ruolo = %s WHERE id_admin = %s",
            (username, ruolo, id_admin)
        )
        connessione.commit()
        cursor.close()
    finally:
        connessione.close()
 
    return RedirectResponse("/admin/utenti?esito=utente-modificato", status_code=303)
 
 
# ---------- Eliminazione utente (solo Amministratore) ----------
 
@router.post("/utenti/{id_admin}/elimina")
def elimina_utente(request: Request, id_admin: int):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect
 
    connessione = get_connection()
    try:
        cursor = connessione.cursor(dictionary=True)
        cursor.execute("SELECT id_admin, username, ruolo FROM Amministratore WHERE id_admin = %s", (id_admin,))
        utente = cursor.fetchone()
        cursor.close()
 
        if not utente:
            raise HTTPException(status_code=404, detail="Utente non trovato")
 
        if utente["username"] == admin["username"]:
            return RedirectResponse("/admin/utenti?esito=errore-autoeliminazione", status_code=303)
 
        if utente["ruolo"] == "Amministratore" and _conta_amministratori(connessione) <= 1:
            return RedirectResponse("/admin/utenti?esito=errore-ultimo-amministratore", status_code=303)
 
        cursor = connessione.cursor()
        cursor.execute("DELETE FROM Amministratore WHERE id_admin = %s", (id_admin,))
        connessione.commit()
        cursor.close()
    finally:
        connessione.close()
 
    return RedirectResponse("/admin/utenti?esito=utente-eliminato", status_code=303)

@router.post("/utenti/{id_admin}/reimposta-password")
def esegui_reimposta_password(
    request: Request,
    id_admin: int,
    password: str = Form(...),
    conferma_password: str = Form(...),
):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect
 
    connessione = get_connection()
    try:
        cursor = connessione.cursor(dictionary=True)
        cursor.execute("SELECT id_admin, username FROM Amministratore WHERE id_admin = %s", (id_admin,))
        utente = cursor.fetchone()
        cursor.close()
 
        if not utente:
            raise HTTPException(status_code=404, detail="Utente non trovato")
 
        errore = None
        if password != conferma_password:
            errore = "Le due password non coincidono."
        elif len(password) < 8:
            errore = "La password deve avere almeno 8 caratteri."
 
        if errore:
            return templates.TemplateResponse(
                "admin_reimposta_password.html",
                {"request": request, "admin": admin, "utente": utente, "errore": errore}
            )
 
        cursor = connessione.cursor()
        cursor.execute(
            "UPDATE Amministratore SET password_hash = %s WHERE id_admin = %s",
            (hash_password(password), id_admin)
        )
        connessione.commit()
        cursor.close()
    finally:
        connessione.close()
 
    return RedirectResponse("/admin/utenti?esito=password-reimpostata", status_code=303)


# ---------- Importazione da Excel (solo Amministratore) ----------

@router.get("/importa-excel")
def form_importa_excel(request: Request):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect

    return templates.TemplateResponse(
        "admin_importa_excel.html",
        {"request": request, "admin": admin, "risultato": None, "errore": None}
    )


@router.post("/importa-excel")
async def esegui_importa_excel(
    request: Request,
    file_excel: UploadFile = File(...),
    modalita: str = Form("aggiungi"),
    conferma_sostituzione: Optional[str] = Form(None),
):
    admin, redirect = _richiedi_amministratore_html(request)
    if redirect:
        return redirect

    if not file_excel.filename or not file_excel.filename.lower().endswith(".xlsx"):
        return templates.TemplateResponse(
            "admin_importa_excel.html",
            {"request": request, "admin": admin, "risultato": None,
             "errore": "Il file deve essere in formato .xlsx (lo stesso formato del file di test)."}
        )

    if modalita == "sostituisci" and not conferma_sostituzione:
        return templates.TemplateResponse(
            "admin_importa_excel.html",
            {"request": request, "admin": admin, "risultato": None,
             "errore": "Per sostituire completamente il catalogo devi spuntare la casella di conferma."}
        )

    contenuto = await file_excel.read()

    try:
        workbook = load_workbook(io.BytesIO(contenuto))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
    except Exception:
        return templates.TemplateResponse(
            "admin_importa_excel.html",
            {"request": request, "admin": admin, "risultato": None,
             "errore": "Impossibile leggere il file. Verifica che sia un Excel .xlsx valido, con la stessa struttura del file di test."}
        )

    righe_elaborate = 0
    errori = []
    materiali_eliminati = None

    connessione = get_connection()
    try:
        if modalita == "sostituisci":
            allegati_da_rimuovere = crud.elimina_tutti_i_materiali(connessione)
            for nome_file in allegati_da_rimuovere:
                _elimina_allegato_pdf(nome_file)
            materiali_eliminati = None  

        for numero_riga, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            row = dict(zip(headers, values))

            if not row.get("Titolo"):
                continue

            try:
                crud.importa_riga_excel(connessione, row)
                righe_elaborate += 1
            except Exception as errore_riga:
                errori.append(f"Riga {numero_riga}: {errore_riga}")
    finally:
        connessione.close()

    return templates.TemplateResponse(
        "admin_importa_excel.html",
        {
            "request": request, "admin": admin, "errore": None,
            "risultato": {
                "righe_elaborate": righe_elaborate, "errori": errori,
                "sostituito": modalita == "sostituisci",
            },
        }
    )