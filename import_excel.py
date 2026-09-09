"""
Importa i dati dal file Excel nel database, evitando duplicati.

Uso: python import_excel.py [nome_file.xlsx]
Se non specifichi un file, usa Excel_test_catalogazione_musicoterapia.xlsx.

La logica di importazione (che evita i duplicati) è condivisa con la
funzione "Importa da Excel" del pannello admin, in app/crud.py.
"""
import sys

from openpyxl import load_workbook
from dotenv import load_dotenv
from os import getenv

import mysql.connector

from app.crud import importa_riga_excel

load_dotenv()

nome_file = sys.argv[1] if len(sys.argv) > 1 else "Excel_test_catalogazione_musicoterapia.xlsx"

try:
    connessione = mysql.connector.connect(
        host=getenv("DB_HOST"),
        user=getenv("DB_USER"),
        password=getenv("DB_PASSWORD"),
        database=getenv("DB_NAME")
    )
    print("Connessione a MySQL riuscita!")
except mysql.connector.Error as err:
    print(f"Errore durante la connessione a MySQL: {err}")
    exit()

workbook = load_workbook(nome_file)
sheet = workbook.active
headers = [cell.value for cell in sheet[1]]

righe_elaborate = 0
righe_saltate = 0

for numero_riga, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
    row = dict(zip(headers, values))

    if not row.get("Titolo"):
        righe_saltate += 1
        continue

    try:
        importa_riga_excel(connessione, row)
        righe_elaborate += 1
    except Exception as errore:
        print(f"Errore alla riga {numero_riga}: {errore}")
        righe_saltate += 1

print(f"Importazione completata. Righe elaborate: {righe_elaborate}, righe saltate: {righe_saltate}")

connessione.close()