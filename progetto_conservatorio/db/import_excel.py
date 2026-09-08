from openpyxl import load_workbook
from dotenv import load_dotenv
from os import getenv

import mysql.connector

load_dotenv()

try:
    connessione = mysql.connector.connect(
        host=getenv("DB_HOST"),
        user=getenv("DB_USER"),
        password=getenv("DB_PASSWORD"),
        database=getenv("DB_NAME")
    )

    print("Connessione a MySQL riuscita!")

except mysql.connector.Error as err:
    print(f"Errore durante la connessione a MySQL: {err}")
    exit()

workbook = load_workbook("Excel_test_catalogazione_musicoterapia.xlsx")
sheet = workbook.active


def get_or_create_editore(connessione, nome, citta):
    cursor = connessione.cursor()

    cursor.execute(
        "SELECT id_editore FROM Editore WHERE nome = %s",
        (nome,)
    )

    risultato = cursor.fetchone()

    if risultato:
        cursor.close()
        return risultato[0]

    cursor.execute(
        "INSERT INTO Editore (nome, citta) VALUES (%s, %s)",
        (nome, citta)
    )

    connessione.commit()
    id_nuovo_editore = cursor.lastrowid
    cursor.close()

    return id_nuovo_editore

def get_or_create_lingua(connessione, nome):
    cursor = connessione.cursor()

    cursor.execute(
        "SELECT id_lingua FROM Lingua WHERE nome = %s",
        (nome,)
    )

    risultato = cursor.fetchone()

    if risultato:
        cursor.close()
        return risultato[0]

    cursor.execute(
        "INSERT INTO Lingua (nome) VALUES (%s)",
        (nome,)
    )

    connessione.commit()
    id_nuova_lingua = cursor.lastrowid
    cursor.close()

    return id_nuova_lingua

def get_or_create_materiale(
    connessione,
    titolo,
    anno,
    num_pagine,
    isbn,
    abstract,
    indice,
    tipologia,
    ebp,
    livello,
    link_catalogo,
    id_editore,
    id_lingua_testo,
    id_lingua_originale
):
    cursor = connessione.cursor()

    if isbn:
        cursor.execute(
            """
            SELECT id_materiale
            FROM Materiale
            WHERE isbn = %s
            LIMIT 1
            """,
            (isbn,)
        )
    else:
        cursor.execute(
            """
            SELECT id_materiale
            FROM Materiale
            WHERE titolo = %s
              AND anno = %s
              AND tipologia = %s
            LIMIT 1
            """,
            (titolo, anno, tipologia)
        )

    risultato = cursor.fetchone()

    if risultato:
        id_materiale = risultato[0]
        cursor.close()
        return id_materiale

    cursor.execute(
        """
        INSERT INTO Materiale (
            titolo,
            anno,
            num_pagine,
            isbn,
            abstract,
            indice,
            tipologia,
            ebp,
            livello,
            link_catalogo,
            id_editore,
            id_lingua_testo,
            id_lingua_originale
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            titolo,
            anno,
            num_pagine,
            isbn,
            abstract,
            indice,
            tipologia,
            ebp,
            livello,
            link_catalogo,
            id_editore,
            id_lingua_testo,
            id_lingua_originale
        )
    )

    connessione.commit()

    id_materiale = cursor.lastrowid
    cursor.close()

    return id_materiale

def get_or_create_copia(
    connessione,
    numero,
    collocazione,
    stato,
    note,
    nota_critica,
    id_materiale
):
    cursor = connessione.cursor()

    cursor.execute(
        """
        SELECT id_copia
        FROM Copia
        WHERE numero = %s
        """,
        (numero,)
    )

    risultato = cursor.fetchone()

    if risultato:
        cursor.close()
        return risultato[0]

    cursor.execute(
        """
        INSERT INTO Copia (
            numero,
            collocazione,
            stato,
            note,
            nota_critica,
            id_materiale
        )
        VALUES (%s, %s, %s, %s, %s, %s)
        """,
        (
            numero,
            collocazione,
            stato,
            note,
            nota_critica,
            id_materiale
        )
    )

    connessione.commit()

    id_copia = cursor.lastrowid
    cursor.close()

    return id_copia

def get_or_create_persona(connessione, nome_completo):
    cursor = connessione.cursor()

    parti = nome_completo.split()

    if len(parti) == 1:
        nome = parti[0]
        cognome = ""
    else:
        nome = " ".join(parti[:-1])
        cognome = parti[-1]
    
    cursor.execute(
        """
        SELECT id_persona
        FROM Persona
        WHERE nome = %s AND cognome = %s
        LIMIT 1
        """,
        (nome, cognome)
    )

    risultato = cursor.fetchone()

    if risultato:
        cursor.close()
        return risultato[0]
    
    cursor.execute(
        """
        INSERT INTO Persona (nome, cognome)
        VALUES (%s, %s)
        """,
        (nome, cognome)
    )   

    connessione.commit()
    id_persona = cursor.lastrowid
    cursor.close()

    return id_persona

def collega_materiale_persona(connessione, id_materiale, id_persona, ruolo):
    cursor = connessione.cursor()

    cursor.execute(
        """
        SELECT id_materiale
        FROM Materiale_Persona
        WHERE id_materiale = %s AND id_persona = %s AND ruolo = %s
        LIMIT 1
        """,
        (id_materiale, id_persona, ruolo)
    )

    risultato = cursor.fetchone()

    if not risultato:
        cursor.execute(
            """
            INSERT INTO Materiale_Persona (id_materiale, id_persona, ruolo)
            VALUES (%s, %s, %s)
            """,
            (id_materiale, id_persona, ruolo)
        )

        connessione.commit()

    cursor.close()

def get_or_create_categoria(connessione, nome):
    cursor = connessione.cursor()

    if nome is None or nome.strip() == "":
        return None

    cursor.execute(
        """
        SELECT id_categoria
        FROM Categoria
        WHERE nome = %s
        LIMIT 1
        """,
        (nome,)
    )

    risultato = cursor.fetchone()

    if risultato:
        cursor.close()
        return risultato[0]
    
    cursor.execute(
        """
        INSERT INTO Categoria (nome)
        VALUES (%s)
        """,
        (nome,)
    )

    connessione.commit()

    id_categoria = cursor.lastrowid
    cursor.close()

    return id_categoria

def collega_materiale_categoria(connessione, id_materiale, id_categoria):
    cursor = connessione.cursor()

    cursor.execute(
        """
        SELECT *
        FROM Materiale_Categoria
        WHERE id_materiale = %s AND id_categoria = %s
        LIMIT 1
        """,
        (id_materiale, id_categoria)
    )

    risultato = cursor.fetchone()

    if not risultato:
        cursor.execute(
            """
            INSERT INTO Materiale_Categoria (id_materiale, id_categoria)
            VALUES (%s, %s)
            """,
            (id_materiale, id_categoria)
        )

        connessione.commit()

    cursor.close()

def get_or_create_tag(connessione, nome):

    if nome is None or nome.strip() == "":
        return None
    
    cursor = connessione.cursor()

    cursor.execute(
        """
        SELECT id_tag
        FROM Tag
        WHERE nome = %s
        LIMIT 1
        """,
        (nome,)
    )

    risultato = cursor.fetchone()

    if risultato:
        cursor.close()
        return risultato[0]
    
    cursor.execute(
        """
        INSERT INTO Tag (nome)
        VALUES (%s)
        """,
        (nome,)
    )

    connessione.commit()

    id_tag = cursor.lastrowid
    cursor.close()

    return id_tag

def collega_materiale_tag(connessione, id_materiale, id_tag):
    cursor = connessione.cursor()

    cursor.execute(
        """
        SELECT *
        FROM Materiale_Tag
        WHERE id_materiale = %s
          AND id_tag = %s
        LIMIT 1
        """,
        (id_materiale, id_tag)
    )

    risultato = cursor.fetchone()

    if not risultato:
        cursor.execute(
            """
            INSERT INTO Materiale_Tag
            (id_materiale, id_tag)
            VALUES (%s, %s)
            """,
            (id_materiale, id_tag)
        )

        connessione.commit()

    cursor.close()

def get_or_create_sh(connessione, termine):

    if termine is None or termine.strip() == "":
        return None

    cursor = connessione.cursor()

    cursor.execute(
        """
        SELECT id_sh
        FROM SH
        WHERE termine = %s
        LIMIT 1
        """,
        (termine,)
    )

    risultato = cursor.fetchone()

    if risultato:
        cursor.close()
        return risultato[0]

    cursor.execute(
        """
        INSERT INTO SH (termine)
        VALUES (%s)
        """,
        (termine,)
    )

    connessione.commit()

    id_sh = cursor.lastrowid
    cursor.close()

    return id_sh

def collega_materiale_sh(connessione, id_materiale, id_sh):
    cursor = connessione.cursor()

    cursor.execute(
        """
        SELECT *
        FROM Materiale_SH
        WHERE id_materiale = %s
          AND id_sh = %s
        LIMIT 1
        """,
        (id_materiale, id_sh)
    )

    risultato = cursor.fetchone()

    if not risultato:
        cursor.execute(
            """
            INSERT INTO Materiale_SH
            (id_materiale, id_sh)
            VALUES (%s, %s)
            """,
            (id_materiale, id_sh)
        )

        connessione.commit()

    cursor.close()

headers = [cell.value for cell in sheet[1]]

righe_elaborate = 0

for values in sheet.iter_rows(min_row=2, values_only=True):
    row = dict(zip(headers, values))

    testo_persone = row["Autore / Curatore"]
    lista_persone = testo_persone.split(";")

    id_editore = get_or_create_editore(
        connessione,
        row["Editore"],
        row["Città editore"]
    )

    id_lingua_testo = get_or_create_lingua(
        connessione,
        row["Lingua testo"]
    )

    id_lingua_originale = get_or_create_lingua(
        connessione,
        row["Lingua originale"]
    )

    id_materiale = get_or_create_materiale(
        connessione,
        row["Titolo"],
        row["Anno"],
        row["N. pagine"],
        row["ISBN"],
        row["Abstract"],
        row["Indice"],
        row["Tipologia"],
        row["EBP"],
        row["Livello"],
        row["Link catalogo ABV"],
        id_editore,
        id_lingua_testo,
        id_lingua_originale
    )

    id_categoria = get_or_create_categoria(
        connessione,
        row["Categoria"]
    )

    if id_categoria is not None:
        collega_materiale_categoria(
            connessione,
            id_materiale,
            id_categoria
        )

    for colonna in ["Tag 1", "Tag 2", "Tag 3", "Tag 4"]:
        id_tag = get_or_create_tag(
            connessione,
            row[colonna]
        )

        if id_tag is not None:
            collega_materiale_tag(
                connessione,
                id_materiale,
                id_tag
            )

    for colonna in ["SH 1", "SH 2", "SH 3"]:
        id_sh = get_or_create_sh(
            connessione,
            row[colonna]
        )

        if id_sh is not None:
            collega_materiale_sh(
                connessione,
                id_materiale,
                id_sh
            )

    get_or_create_copia(
        connessione,
        row["N."],
        row["Collocazione"],
        row["Stato"],
        row["Note"],
        row["Nota critica"],
        id_materiale
    )

    for persona in lista_persone:
        persona = persona.strip()

        if not persona:
            continue

        if "(" in persona and persona.endswith(")"):
            nome_completo, ruolo = persona.rsplit("(", 1)
            nome_completo = nome_completo.strip()
            ruolo = ruolo[:-1].strip()
        else:
            nome_completo = persona
            ruolo = "Autore"

        id_persona = get_or_create_persona(
            connessione,
            nome_completo
        )

        collega_materiale_persona(
            connessione,
            id_materiale,
            id_persona,
            ruolo
        )

    righe_elaborate += 1

print("Importazione completata. Righe elaborate: {righe_elaborate}")

connessione.close()